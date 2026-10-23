"""把 docs/FinalReport/畢業專題報告.md 轉成符合 #148 排版規範的 Word 檔。

用法（macOS，需裝 Word 以取得 mathml2omml.xsl）：
uv run --no-project --with python-docx --with latex2mathml --with lxml --with pillow python scripts/opt/report_md2docx.py \
    docs/FinalReport/畢業專題報告.md docs/FinalReport/畢業專題報告.docx \
    "/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl"
產出後在 Word 全選按 F9 更新目錄、圖目錄與表目錄。追蹤連結、> 排版指示與各章寫作提示不會放進 Word。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree

MD, OUT, XSL = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
BASE = MD.parent
EAST = "標楷體"
WEST = "Times New Roman"
MONO = "Consolas"
TRANSFORM = etree.XSLT(etree.parse(str(XSL)))

# 各章開頭的寫作提示與骨架說明，排進 Word 時刪除
GUIDANCE = {
    "封面由上到下為學校與系所、題目、指導教授、學生、日期。正式 Word 全部置中、單行間距。",
    "摘要是全文縮影。最後的關鍵詞使用粗體並縮排。",
    "章名置中，節名靠左，編號與名稱中間空一格。緒論引導讀者進入問題。",
    "回顧相關方法並提出比較對象。選材以學術價值足夠、且與主題、架構或方法有關者為限。",
    "本章呈現本專題的方法，是全文核心。",
    "條目使用正文實際引用的文獻。",
    "附錄使用中文編號：附錄 一、附錄 二、附錄 三。",
}


# ---------------------------------------------------------------- 樣式

def set_fonts(rpr_owner, east=EAST, west=WEST):
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(key), west)
    fonts.set(qn("w:eastAsia"), east)
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(key), None)


def style_setup(doc):
    st = doc.styles
    normal = st["Normal"]
    normal.font.size = Pt(12)
    set_fonts(normal.element.get_or_add_rPr())
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.line_spacing = Pt(16)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def heading(name, size, align, before, after, line=None):
        s = st[name]
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = RGBColor(0, 0, 0)
        set_fonts(s.element.get_or_add_rPr())
        f = s.paragraph_format
        f.alignment = align
        f.space_before = Pt(before)
        f.space_after = Pt(after)
        f.keep_with_next = True
        if line == "1.5":
            f.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        else:
            f.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            f.line_spacing = Pt(16)

    heading("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 12, 12, "1.5")
    heading("Heading 2", 14, WD_ALIGN_PARAGRAPH.JUSTIFY, 12, 6)
    heading("Heading 3", 12, WD_ALIGN_PARAGRAPH.JUSTIFY, 12, 6)
    heading("Heading 4", 12, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 3)

    from docx.enum.style import WD_STYLE_TYPE
    for name, align in (("圖名", WD_ALIGN_PARAGRAPH.CENTER), ("表名", WD_ALIGN_PARAGRAPH.LEFT)):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        s.paragraph_format.alignment = align
        s.paragraph_format.space_before = Pt(6)
        s.paragraph_format.space_after = Pt(6)
        if name == "表名":
            s.paragraph_format.keep_with_next = True
    code = st.add_style("程式碼", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    code.font.size = Pt(9.5)
    set_fonts(code.element.get_or_add_rPr(), EAST, MONO)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for name in ("List Bullet", "List Number"):
        set_fonts(st[name].element.get_or_add_rPr())
        st[name].font.size = Pt(12)


def page_setup(section, fmt=None, start=None, numbered=True):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(3.17)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    sect = section._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    if fmt:
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        cols = sect.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(pg)   # 依 schema，pgNumType 要在 cols 之前
        else:
            sect.append(pg)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    if numbered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(p, "PAGE")


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
            run._r.append(el)
        elif kind == "text":
            if text:
                t = OxmlElement("w:t")
                t.text = text
                run._r.append(t)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")
            run._r.append(el)


# ---------------------------------------------------------------- 公式

def clean_latex(tex):
    tex = tex.replace(r"\dfrac", r"\frac").replace(r"\tfrac", r"\frac")
    tex = re.sub(r"\\\\\[[0-9.]+ex\]", r"\\\\", tex)
    tex = tex.replace(r"\big(", "(").replace(r"\big)", ")").replace(r"\Big(", "(").replace(r"\Big)", ")")
    tex = re.sub(r"\\operatorname\{([^}]*)\}", r"\\mathrm{\1}", tex)
    return tex


def omml(tex):
    mathml = latex_to_mathml(clean_latex(tex))
    tree = etree.fromstring(mathml.encode("utf-8"))
    result = TRANSFORM(tree)
    return result.getroot()


def append_math(paragraph, tex, display=False):
    try:
        el = omml(tex)
    except Exception:
        r = paragraph.add_run(tex)
        set_fonts(r._r.get_or_add_rPr())
        return False
    if display:
        para = OxmlElement("m:oMathPara")
        para.append(el)
        paragraph._p.append(para)
    else:
        paragraph._p.append(el)
    return True


# ---------------------------------------------------------------- 行內格式

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\$[^$]+\$|!\[[^\]]*\]\([^)]*\)|\[[^\]]+\]\([^)]*\))")


def add_inline(paragraph, text, bold=False):
    text = re.sub(r"<a id=\"[^\"]+\"></a>", "", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(paragraph, part[2:-2], bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_fonts(r._r.get_or_add_rPr(), EAST, MONO)
            r.font.size = Pt(11)
            r.bold = bold
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            append_math(paragraph, part[1:-1])
        elif re.fullmatch(r"\[([^\]]+)\]\(([^)]*)\)", part):
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]*)\)", part)
            label, target = m.group(1), m.group(2)
            if target.startswith("#ref-"):
                label = f"[{label}]"
            add_inline(paragraph, label, bold)
        else:
            r = paragraph.add_run(part)
            r.bold = bold


# ---------------------------------------------------------------- 區塊

def table_from(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # 欄寬依各欄內容長度分配（上限避免單一長欄吃掉整表），總寬為版心 14.66 cm
    def weight(text):
        text = re.sub(r"[*`$]", "", text)
        return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
    weights = []
    for j in range(ncol):
        col = [r[j] if j < len(r) else "" for r in [header] + body]
        longest = max(weight(c) for c in col)
        avg = sum(weight(c) for c in col) / len(col)
        weights.append(max(6, min(60, 0.5 * longest + 0.5 * avg)))
    total = sum(weights)
    widths = [Cm(max(1.3, 14.66 * w / total)) for w in weights]
    scale = Cm(14.66) / sum(widths)
    widths = [int(w * scale) for w in widths]
    font = 8 if ncol >= 8 else 9 if ncol >= 5 else 10
    for i, row in enumerate([header] + body):
        for j in range(ncol):
            val = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            cell.width = widths[j]
            add_inline(p, val, bold=(i == 0))
            for r in p.runs:
                r.font.size = Pt(font)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    doc.add_paragraph()


def add_image(doc, alt, path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    from PIL import Image as _Im
    w, h = _Im.open(BASE / path).size
    width = Cm(14.5)
    if h / w * 14.5 > 20.0:   # 太長的圖以高度 20 cm 為上限，圖名才能與圖同頁
        width = Cm(20.0 * w / h)
    p.add_run().add_picture(str(BASE / path), width=width)


def display_equation(doc, tex):
    tex = tex.strip()
    m = re.search(r"\\qquad\s*\((\d+)\)\s*$", tex)
    number = None
    if m:
        number, tex = m.group(1), tex[:m.start()].strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ok = append_math(p, tex, display=False)
    if number:
        p.paragraph_format.tab_stops.add_tab_stop(Cm(14.66), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(7.33), WD_TAB_ALIGNMENT.CENTER)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # 置中公式＋右側編號：以兩個 tab 對齊
        first = OxmlElement("w:r")
        tab = OxmlElement("w:tab")
        first.append(tab)
        p._p.insert(len(p._p.findall(qn("w:pPr"))), first)
        r = p.add_run()
        r.add_tab()
        r2 = p.add_run(f"({number})")
    return ok


# ---------------------------------------------------------------- 主流程

def build():
    lines = MD.read_text(encoding="utf-8").split("\n")
    start = next(i for i, l in enumerate(lines) if l.strip() == "## 封面")
    lines = lines[start:]

    doc = Document()
    style_setup(doc)
    page_setup(doc.sections[0], numbered=False)

    # 封面
    def cover_line(text, size, bold=False, before=0, indent=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if indent is None else WD_ALIGN_PARAGRAPH.LEFT
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        return p

    cover = {}
    sect = None
    i = 0
    while i < len(lines) and not lines[i].startswith("## 誌謝"):
        l = lines[i]
        if l.startswith("### "):
            sect = l[4:].strip()
            cover[sect] = []
        elif sect and l.strip() and not l.startswith(">") and not l.startswith("追蹤："):
            cover[sect].append(l.rstrip())
        i += 1
    for _ in range(4):
        doc.add_paragraph()
    for t in cover["學校、系所與報告類型"]:
        cover_line(t, 32, before=3)
    for _ in range(3):
        doc.add_paragraph()
    for t in cover["專題題目"]:
        cover_line(t, 24, bold=True)
    for _ in range(4):
        doc.add_paragraph()
    for t in cover["指導教授"]:
        cover_line(t, 22)
    doc.add_paragraph()
    for t in cover["學生"]:
        cover_line(t, 22, indent=3.6)
    for _ in range(3):
        doc.add_paragraph()
    for t in cover["日期"]:
        cover_line(t, 20)

    lines = lines[i:]
    chapter = None
    started_body = False
    in_refs = False
    skip_section = False
    i = 0
    while i < len(lines):
        l = lines[i]
        s = l.strip()

        if l.startswith("## "):
            title = l[3:].strip()
            chapter = title
            skip_section = False
            in_refs = title == "參考文獻"
            if title == "誌謝":
                page_setup(doc.add_section(WD_SECTION.NEW_PAGE), "lowerRoman", 1)
            elif title.startswith("第一章") and not started_body:
                page_setup(doc.add_section(WD_SECTION.NEW_PAGE), "decimal", 1)
                started_body = True
            else:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=1)
            add_inline(h, title)
            if title == "目錄":
                add_field(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u', "（在 Word 按 F9 更新目錄）")
                skip_section = True
            elif title == "圖目錄":
                add_field(doc.add_paragraph(), r'TOC \h \z \t "圖名,1"', "（在 Word 按 F9 更新圖目錄）")
                skip_section = True
            elif title == "表目錄":
                add_field(doc.add_paragraph(), r'TOC \h \z \t "表名,1"', "（在 Word 按 F9 更新表目錄）")
                skip_section = True
            i += 1
            continue

        if skip_section:
            i += 1
            continue

        if l.startswith("### ") or l.startswith("#### "):
            level = 2 if l.startswith("### ") else 3
            title = l.lstrip("#").strip()
            if chapter == "摘要":
                i += 1
                continue
            if in_refs:
                skip_section = title != "全域編號書目"
                i += 1
                continue
            h = doc.add_heading(level=level)
            add_inline(h, title)
            i += 1
            continue

        if not s or s.startswith("追蹤：") or s.startswith(">") or s in GUIDANCE:
            i += 1
            continue

        if s.startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                p = doc.add_paragraph(style="程式碼")
                r = p.add_run(lines[j])
                i = j
                j += 1
            i = j + 1
            continue

        if s == "$$":
            j = i + 1
            buf = []
            while lines[j].strip() != "$$":
                buf.append(lines[j])
                j += 1
            display_equation(doc, "\n".join(buf))
            i = j + 1
            continue

        if s.startswith("|"):
            j = i
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(lines[j])
                j += 1
            table_from(doc, rows)
            i = j
            continue

        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            add_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        if re.match(r"^圖\d+\.\d+\s", s) and len(s) <= 40 and "。" not in s:
            p = doc.add_paragraph(style="圖名")
            add_inline(p, s)
            i += 1
            continue
        if re.match(r"^表\d+\.\d+\s", s) and len(s) <= 60 and "。" not in s:
            p = doc.add_paragraph(style="表名")
            add_inline(p, s)
            i += 1
            continue

        m = re.match(r"^(\s*)- (.*)$", l)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\. (.*)$", l)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.85)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, f"{m.group(2)}.\u3000" + m.group(3))
            i += 1
            continue

        p = doc.add_paragraph()
        plain = re.sub(r"[*`$]", "", re.sub(r"\]\([^)]*\)", "]", s))
        if max((len(t) for t in re.findall(r"[\x21-\x7e]+", plain)), default=0) > 22:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT   # 長路徑或英文無法斷行，左右對齊會把同行中文字距拉開
        if in_refs:
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.85)
        elif chapter == "摘要" and s.startswith("\u3000\u3000**關鍵詞"):
            add_inline(p, s)
            i += 1
            continue
        else:
            p.paragraph_format.first_line_indent = Pt(24)
        add_inline(p, s)
        i += 1

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(OUT)


if __name__ == "__main__":
    build()
